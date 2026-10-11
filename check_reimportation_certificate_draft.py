# Publication example: personal paths and shipment outputs removed.
# Run in IPython/Jupyter for styled table display.
import pandas as pd
from openpyxl import load_workbook
from IPython.display import display

# ============================================================
# PART 1: อ่าน Invoice Excel -> df_inv
# ============================================================
file_path_inv   = r"data/invoice.xlsx"
file_path_draft = r"data/draft.xlsx"  # เปลี่ยนเป็น path จริง

wb = load_workbook(file_path_inv, read_only=True)
ws = wb['INVOICE']

invoice_date = None
invoice_no   = None

for i, row in enumerate(ws.iter_rows(max_row=30, values_only=True), 1):
    for cell in row:
        if hasattr(cell, 'year') and invoice_date is None:
            invoice_date = cell
    if i == 4:
        invoice_no = row[34]

data_rows = []
name_map  = {}

for i, row in enumerate(ws.iter_rows(min_row=32, values_only=True), 32):
    vals = {j: v for j, v in enumerate(row) if v is not None}
    if 22 in vals and isinstance(vals[22], (int, float)):
        data_rows.append({
            'excel_row'       : i,
            'Model No.'       : vals.get(6, '').split('/')[0].strip() if '/' in str(vals.get(6, '')) else vals.get(6, ''),
            'HS Code'         : vals.get(6, '').split('/')[1].strip() if '/' in str(vals.get(6, '')) else '',
            "Q'ty"            : vals.get(22),
            'Unit'            : vals.get(25),
            'Unit Price (JPY)': vals.get(28),
            'S.Total (JPY)'   : vals.get(33),
            'Origin'          : vals.get(38),
        })
    elif 6 in vals and 22 not in vals and i > 33:
        raw = str(vals[6]).strip()
        if raw and raw != '\u3000' and raw != 'MODEL NAME':
            name_map[i] = raw.replace('\n', ' ')

invoice_date_str = invoice_date.strftime('%d/%m/%Y') if invoice_date else ''

for d in data_rows:
    next_row = d['excel_row'] + 1
    d['Model Name'] = name_map.get(next_row, '')
    del d['excel_row']

df_inv = pd.DataFrame(data_rows, columns=[
    'Model No.', 'HS Code', 'Model Name',
    "Q'ty", 'Unit', 'Unit Price (JPY)', 'S.Total (JPY)', 'Origin'
])
df_inv.insert(0, 'Invoice No.', invoice_no)
df_inv.insert(1, 'Date', invoice_date_str)
df_inv['Model No.'] = df_inv['Model No.'].astype(str).str.strip()

# ============================================================
# PART 2: อ่าน df_draft (ใบขน) จากไฟล์ Excel
# ============================================================
df_draft = pd.read_excel(file_path_draft, engine='openpyxl')
df_draft['พิกัด'] = df_draft['พิกัด'].apply(
    lambda x: str(int(x)).zfill(8) if pd.notna(x) else ''
)
df_draft['รหัสสินค้า'] = df_draft['รหัสสินค้า'].astype(str).str.strip()

# ============================================================
# PART 3: df_master (ข้อมูลหลัก)
# ============================================================
# FICTIONAL MASTER ONLY. XX and tariff codes are placeholders, not valid customs data.
rows = [{'Tariff Code': '000000000001', 'Part No': 'DEMO-PART-001', 'Thai Description': 'บรรจุภัณฑ์ตัวอย่าง 1', 'English Description': 'DEMO RETURNABLE PACKAGE 1', 'Origin Country': 'XX', 'Net Weight': 1.0, 'Currency Code': 'JPY', 'Qty Unit': 'C62', 'BRAND': 'DEMO BRAND'}, {'Tariff Code': '000000000002', 'Part No': 'DEMO-PART-002', 'Thai Description': 'บรรจุภัณฑ์ตัวอย่าง 2', 'English Description': 'DEMO RETURNABLE PACKAGE 2', 'Origin Country': 'XX', 'Net Weight': 2.0, 'Currency Code': 'JPY', 'Qty Unit': 'C62', 'BRAND': 'DEMO BRAND'}, {'Tariff Code': '000000000003', 'Part No': 'DEMO-PART-003', 'Thai Description': 'บรรจุภัณฑ์ตัวอย่าง 3', 'English Description': 'DEMO RETURNABLE PACKAGE 3', 'Origin Country': 'XX', 'Net Weight': 3.0, 'Currency Code': 'JPY', 'Qty Unit': 'C62', 'BRAND': 'DEMO BRAND'}]

df_master = pd.DataFrame(rows)
df_master["Tariff Code"] = df_master["Tariff Code"].astype(str)
df_master["Part No"]     = df_master["Part No"].astype(str).str.strip()
df_master["Net Weight"]  = pd.to_numeric(df_master["Net Weight"], errors="coerce")

# set index เพื่อ lookup เหมือน VLOOKUP
master_lookup = df_master.set_index('Part No')
inv_lookup    = df_inv.set_index('Model No.')

# ============================================================
# PART 4: เช็คข้อมูล โดย match จาก Part No (VLOOKUP style)
# ============================================================
unit_equivalents = {'c62': 'pc', 'pc': 'c62'}

def normalize(val):
    s = str(val).strip().lower().replace(' ', '')
    try:
        s = str(float(s))
    except:
        pass
    return s

def is_equal(cmp_a, cmp_b, field):
    if cmp_a == cmp_b:
        return True
    if field == 'หน่วยปริมาณ (Invoice)':
        if unit_equivalents.get(cmp_a) == cmp_b:
            return True
        if unit_equivalents.get(cmp_b) == cmp_a:
            return True
    return False

col_mapping = {
    'เลขที่อินวอยซ์'             : ('เลขที่อินวอยซ์',              'Invoice No.',          'inv'),
    'รหัสสินค้า'                  : ('รหัสสินค้า',                   'Part No',              'master'),
    'Brand'                       : ('Brand',                        'BRAND',                'master'),
    'ชื่อสินค้าภาษาอังกฤษ 1'     : ('ชื่อสินค้าภาษาอังกฤษ 1',     'English Description',  'master'),
    'ชื่อสินค้าภาษาไทย 1'        : ('ชื่อสินค้าภาษาไทย 1',        'Thai Description',     'master'),
    'พิกัด'                       : ('พิกัด',                        'Tariff Code',          'master'),
    'ปริมาณ (Invoice)'            : ('ปริมาณ (Invoice)',             "Q'ty",                 'inv'),
    'น้ำหนักสุทธิ'                : ('น้ำหนักสุทธิ',                 'Net Weight * Qty',     'calc'),
    'Unit price (เงินต่างประเทศ)' : ('Unit price (เงินต่างประเทศ)', 'Unit Price (JPY)',     'inv'),
    'หน่วยปริมาณ (Invoice)'       : ('หน่วยปริมาณ (Invoice)',        'Qty Unit',             'master'),
    'รหัสสกุลเงิน'                : ('รหัสสกุลเงิน',                 'Currency Code',        'master'),
    'รหัสประเทศปลายทาง'           : ('รหัสประเทศปลายทาง',            'Origin Country',       'master'),
}

results = []

for _, row_draft in df_draft.iterrows():
    part_no = str(row_draft.get('รหัสสินค้า', '')).strip()

    # --- VLOOKUP จาก Part No ---
    if part_no not in master_lookup.index:
        results.append({
            'Part No'      : part_no,
            'Field'        : '-',
            'df_draft'     : '-',
            'df_master/inv': '-',
            'Status'       : '❌ ไม่พบใน Master',
        })
        continue

    if part_no not in inv_lookup.index:
        results.append({
            'Part No'      : part_no,
            'Field'        : '-',
            'df_draft'     : '-',
            'df_master/inv': '-',
            'Status'       : '❌ ไม่พบใน Invoice',
        })
        continue

    row_master = master_lookup.loc[part_no]
    row_inv    = inv_lookup.loc[part_no]
    qty        = row_inv.get("Q'ty", 1)
    net_w      = row_master.get('Net Weight', 0)

    for field, (draft_col, other_col, source) in col_mapping.items():
        val_draft = str(row_draft.get(draft_col, '')).strip()

        if source == 'master':
            val_other = str(row_master.get(other_col, '')).strip()
        elif source == 'inv':
            val_other = str(row_inv.get(other_col, '')).strip()
        elif source == 'calc':
            calc      = net_w * qty
            val_other = str(round(float(calc), 1))
            val_draft = str(round(float(val_draft), 1)) if val_draft not in ('', 'nan') else ''

        cmp_draft = normalize(val_draft)
        cmp_other = normalize(val_other)

        if cmp_draft == '' or cmp_other == '':
            status = 'ℹ️'
        else:
            status = '✅' if is_equal(cmp_draft, cmp_other, field) else '⚠️ ไม่ตรง'

        results.append({
            'Part No'      : part_no,
            'Field'        : field,
            'df_draft'     : val_draft,
            'df_master/inv': val_other,
            'Status'       : status,
        })

if not results:
    print("⚠️ ไม่มีข้อมูลให้เช็ค")
else:
    df_result   = pd.DataFrame(results)
    df_mismatch = df_result[df_result['Status'] == '⚠️ ไม่ตรง'].reset_index(drop=True)

    pd.set_option('display.max_rows', None)
    pd.set_option('display.max_columns', None)
    pd.set_option('display.max_colwidth', None)

    print("=" * 80)
    print(f"  ผลการตรวจสอบ Invoice: {invoice_no}  วันที่: {invoice_date_str}")
    print("=" * 80)

    if df_mismatch.empty:
        print("✅ ข้อมูลทุก field ตรงกันหมด!")
    else:
        print(f"⚠️ พบข้อมูลไม่ตรงกัน {len(df_mismatch)} รายการ\n")
        display(df_mismatch.style
            .applymap(lambda v: 'background-color: #ffe0e0' if v == '⚠️ ไม่ตรง' else
                                'background-color: #ffd0d0' if '❌' in str(v) else '')
            .set_properties(**{'text-align': 'left'})
            .set_table_styles([
                {'selector': 'th', 'props': [('background-color', '#4472C4'), ('color', 'white'), ('font-weight', 'bold')]},
                {'selector': 'tr:nth-child(even)', 'props': [('background-color', '#DEEAF1')]},
            ])
        )